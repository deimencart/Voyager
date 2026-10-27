import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

COLOR_PRIMARY = RGBColor(31, 73, 125)    # Navy Blue
COLOR_SECONDARY = RGBColor(89, 89, 89)  # Dark Gray
COLOR_TEXT = RGBColor(38, 38, 38)        # Off-black
COLOR_ALERT = RGBColor(180, 40, 40)      # Crimson Red

def set_para_font(p, name='Calibri', size=11, color=COLOR_TEXT, bold=False, italic=False):
    for run in p.runs:
        run.font.name = name
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.bold = bold
        run.italic = italic

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    if level == 1:
        set_para_font(h, 'Calibri', 18, COLOR_PRIMARY, bold=True)
    elif level == 2:
        set_para_font(h, 'Calibri', 14, COLOR_PRIMARY, bold=True)
    elif level == 3:
        set_para_font(h, 'Calibri', 12, COLOR_SECONDARY, bold=True)
    return h

def format_table(tbl, col_widths, headers, data):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_para_font(p, 'Calibri', 10, RGBColor(255, 255, 255), bold=True)
        shd = parse_xml(r'<w:shd {} w:fill="1F497D"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)

    for r_idx, row_data in enumerate(data):
        row_cells = tbl.add_row().cells
        fill_color = "F2F5F8" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_para_font(p, 'Calibri', 9.5, COLOR_TEXT)
            shd = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), fill_color))
            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

    for row in tbl.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

def add_callout(doc, text, title="ADVERTENCIA DE SEGURIDAD"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.4)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.color.rgb = COLOR_ALERT
    r_text = p.add_run(text)
    r_text.font.size = Pt(10)
    r_text.font.italic = True

def build_docx():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Title
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = t.add_run("BITÁCORA TÉCNICA DE PRUEBAS — ROVER VOYAGER 3\n")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = COLOR_PRIMARY
    r_title.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub.add_run("Sesión de Verificación Inalámbrica, Telemetría QGC/Mission Planner, Sensores y Auditoría Eléctrica LiDAR SF40/C\nFecha: 24 de Octubre de 2026")
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    add_styled_heading(doc, "1. Resumen Ejecutivo de la Sesión", 1)
    p = doc.add_paragraph()
    p.add_run(
        "Durante la sesión del 24 de octubre de 2026 se logró consolidar la conexión completamente inalámbrica del rover Voyager 3 "
        "con la computadora portátil mediante la zona de cobertura WiFi DroneLab (192.168.137.153) sin cables Ethernet ni bloqueos de puerta de enlace. "
        "Se verificó la compatibilidad y lectura en tiempo real con QGroundControl v5.1.5 y Mission Planner 1.3.x, validando la triple redundancia inercial (IMU), "
        "la sincronización UART del GPS u-blox Neo-M8N, y el carril de servos a 5.57V. Asimismo, se completó una auditoría eléctrica crítica para la integración "
        "del LiDAR rotativo 360° LightWare SF40/C, desmintiendo la creencia de alimentación a 5V y estableciendo la obligatoriedad de alimentación externa a 12V DC."
    )

    add_styled_heading(doc, "2. Arquitectura de Red y Topología de Operación", 1)
    p = doc.add_paragraph()
    p.add_run("Se resolvieron las incógnitas sobre la topología de red y los servicios heredados del proyecto europeo RAWFIE:")
    bullets_red = [
        ("WiFi DroneLab (Configuración Activa): ", "La Raspberry Pi 3B se asocia de forma totalmente automática a la zona móvil de Windows (192.168.137.153). Ofrece aislamiento limpio, baja latencia (3-8 ms) y compatibilidad total con MAVLink UDP 14550."),
        ("Dilema Swarm vs AP Autónomo: ", "El modo actual de cliente en una red unificada permite escalabilidad a enjambres (Swarms) con comunicación Robot-a-Robot (V2V) y ROS 2 DDS. El modo AP en el router MikroTik (SSID: Voyager_3) aísla el robot en su propia subred privada."),
        ("Dicotomía Corto Alcance vs Misiones Largas (GSM / BVLOS): ", "Para operaciones locales (< 80 m), el enlace es 100% WiFi directo sin SIM, sin GSM y sin requerir Internet en la laptop (usando Hotspot o la red Voyager_3 del robot). Para misiones largas (> 100 m hasta kilómetros), el WiFi se atenúa y SÍ es necesario el módem celular GSM/4G con tarjeta SIM conectado a Internet (o radios de telemetría RF SiK 433/915 MHz en TELEM 1)."),
        ("Desmitificación del Error 'RPI is OFFLINE': ", "El script runDrone.sh arrojaba la alerta 'RPI is OFFLINE please check Mikrotik' debido a un ping fallido a un servidor universitario en Grecia (195.134.71.99) del proyecto RAWFIE de 2018. Se verificó que este fallo es irrelevante para el control local.")
    ]
    for b_title, b_desc in bullets_red:
        bp = doc.add_paragraph(style='List Bullet')
        r1 = bp.add_run(b_title)
        r1.bold = True
        bp.add_run(b_desc)

    add_styled_heading(doc, "3. Verificación de Telemetría y Sensores", 1)
    headers_sens = ["Subsistema", "Parámetro / MAVLink", "Estado Observado", "Diagnóstico Físico"]
    widths_sens = [1.8, 1.8, 1.4, 2.5]
    data_sens = [
        ["Identificación Autopiloto", "type: 10, autopilot: 3", "Active (Manual / Hold)", "Pixhawk 1 con ArduRover v3.4.2 plenamente comunicante."],
        ["IMU Triple Redundante", "gx, gy (1, 2, 3) / roll", "roll: -0.59° dinámico", "Giroscopios y acelerómetros redundantes activos sin deriva."],
        ["Receptor GNSS (GPS)", "gpsstatus: 1, satcount: 0", "No Fix (Interior)", "Módulo u-blox Neo-M8N detectado y comunicando por UART en SERIAL 3."],
        ["Alimentación de Servos", "servovoltage: 5576 mV", "5.57 V Activo", "El BEC del variador electrónico alimenta correctamente el rail MAIN OUT."],
        ["Módem Celular 4G", "Huawei USB Dongle", "Sin SIM (Desconectado)", "No afecta al GPS ni al control local; solo brindaba VPN remota."]
    ]
    t_sens = doc.add_table(rows=1, cols=4)
    format_table(t_sens, widths_sens, headers_sens, data_sens)

    add_styled_heading(doc, "4. Auditoría Eléctrica y Conexionado del LiDAR LightWare SF40/C", 1)
    p_lidar = doc.add_paragraph()
    p_lidar.add_run(
        "Se realizó un análisis de primeros principios y cotejo de datasheets oficiales para prevenir la destrucción del sensor o del autopiloto:"
    )
    add_callout(doc, 
        "PROHIBIDO alimentar el LiDAR LightWare SF40/C desde el pin de 5V de la Pixhawk o la Raspberry Pi. "
        "El motor interno de rotación requiere un rango de 6.5V a 30V DC (nominal 12V DC) con picos de 0.5A a 1.0A. "
        "Conectarlo a 5V provoca caída catastrófica de tensión (brownout) y apagón de la electrónica de control. "
        "Conectarlo a los 12V de la placa de distribución (PDB) con masa común y señales lógicas a SERIAL 4/5.",
        "CRÍTICO: ESPECIFICACIÓN ELÉCTRICA"
    )

    headers_lidar = ["Línea", "Tensión Requerida", "Punto de Conexión en Rover", "Notas de Seguridad"]
    widths_lidar = [1.5, 1.6, 2.0, 2.4]
    data_lidar = [
        ["Alimentación VCC (+)", "12.0 V DC (6.5V - 30V)", "Regleta / Salida 12V de la PDB", "Cable rojo. Fusible / BEC dedicado 1A."],
        ["Masa Común (GND)", "0.0 V (Referencia)", "GND de la PDB y GND Pixhawk", "Cable negro. Obligatorio compartir referencia con Pixhawk."],
        ["Lógica Serie (TX)", "3.3 V TTL", "SERIAL 4/5 Pin RX (Pixhawk)", "Transmite nube de puntos de proximidad a 115.200 bps."],
        ["Lógica Serie (RX)", "3.3 V TTL", "SERIAL 4/5 Pin TX (Pixhawk)", "Recibe comandos de configuración de ArduRover."],
        ["Pin 5V Pixhawk", "NO CONECTAR (Dejar libre)", "Puerto SERIAL 4/5 de la Pixhawk", "Si entran 12V a este pin se destruye la Pixhawk."]
    ]
    t_lidar = doc.add_table(rows=1, cols=4)
    format_table(t_lidar, widths_lidar, headers_lidar, data_lidar)

    add_styled_heading(doc, "5. Subsistema de Cámara y Servos Pan/Tilt", 1)
    p_cam = doc.add_paragraph()
    p_cam.add_run(
        "Se definieron las vías para transmitir vídeo en vivo hacia la estación de tierra (laptop) sin saturar la Raspberry Pi:\n"
        "• Opción QGroundControl: Flujo UDP RTP H.264 en el puerto 5600 mediante pipeline de raspivid y GStreamer.\n"
        "• Opción Mission Planner: Inserción directa en el HUD de vuelo mediante flujo HTTP MJPEG a 30 fps.\n"
        "• Servos de Cámara (Pan/Tilt): El soporte cuenta con servos de orientación gobernables desde la Pixhawk mediante parámetros MNT_ (Set Mount) o desde pines auxiliares."
    )

    add_styled_heading(doc, "6. Protocolo de Apagado Seguro Realizado", 1)
    p_off = doc.add_paragraph()
    p_off.add_run(
        "Para preservar la integridad de la tarjeta microSD y prevenir corrupción de sectores en Linux:\n"
        "1. Ejecución de sudo poweroff por consola SSH (192.168.137.153).\n"
        "2. Espera de 10 segundos hasta extinción del LED verde de actividad.\n"
        "3. Desconexión física de la alimentación de potencia de 12V y batería del robot."
    )

    out_path = r"D:\Voyager\analisis_sistema\BITACORA_SESION_2026-10-24.docx"
    doc.save(out_path)
    print(f"Documento generado exitosamente en: {out_path}")

if __name__ == '__main__':
    build_docx()
